"""Small, deterministic fixtures. They contain fictional data only."""
from io import BytesIO
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont


def control(doc, tag, text, locked=False):
    p = doc.add_paragraph(text)
    sdt = OxmlElement("w:sdt")
    props = OxmlElement("w:sdtPr")
    for name, value in [("tag", tag), ("alias", tag)]:
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:val"), value)
        props.append(node)
    if locked:
        node = OxmlElement("w:lock")
        node.set(qn("w:val"), "sdtContentLocked")
        props.append(node)
    content = OxmlElement("w:sdtContent")
    p._p.addprevious(sdt)
    content.append(p._p)
    sdt.append(props)
    sdt.append(content)


def sample_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    doc.add_heading("投标助手 · 功能验证样本", 0)
    doc.add_paragraph("仅用于功能测试 · 所有项目、人员及证书信息均为虚构。")
    doc.add_heading("项目基本信息", 1)
    control(doc, "field:project_name", "【待填充项目名称】")
    control(doc, "field:project_number", "【待填充项目编号】")
    control(doc, "fixed:commitment", "固定承诺：我方确认已阅读招标文件，并对所提交材料的真实性负责。", True)
    doc.add_heading("目录", 1)
    p = doc.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    p._p.append(field)
    doc.add_page_break()
    doc.add_heading("第一章 商务响应", 1)
    control(doc, "section:business", "商务响应测试段落。服务期限为 365 天，响应招标文件规定的交付安排。")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["条款编号", "响应内容", "说明"]):
        cell.text = text
    for values in [("REQ-001", "365 天", "服务期限"), ("REQ-002", "待补充证明材料", "企业资质")]:
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    doc.add_heading("第二章 技术方案", 1)
    control(doc, "section:technical", "技术方案测试段落。建立分阶段实施计划，明确交付清单与质量检查流程。")
    p = doc.add_paragraph("混合格式测试：")
    p.add_run("重点内容").bold = True
    p.add_run("，以及 ")
    p.add_run("保留的强调文字").italic = True
    doc.add_heading("第三章 证明材料", 1)
    control(doc, "asset:certificate", "【证书图片插入区】")
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run("功能验证样本 · 第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    footer.add_run(" 页")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def certificate_png():
    im = Image.new("RGB", (1000, 670), "#f6f3e9")
    d = ImageDraw.Draw(im)
    d.rectangle((28, 28, 970, 640), outline="#284e48", width=4)
    d.rectangle((44, 44, 954, 624), outline="#aa9263", width=2)
    try:
        font_path = "C:/Windows/Fonts/msyh.ttc"
        title = ImageFont.truetype(font_path, 44)
        body = ImageFont.truetype(font_path, 26)
    except OSError:
        title = body = ImageFont.load_default()
    for y, text, font in [(125, "DEMO / 测试证书", title), (245, "企业名称：示例工程技术有限公司", body),
                           (310, "证书编号：DEMO-2026-001", body), (375, "有效期限：2026.01.01 — 2027.12.31", body),
                           (520, "虚构资料 · 仅供文档功能验证使用", body)]:
        d.text((95, y), text, font=font, fill="#284e48")
    out = BytesIO()
    im.save(out, "PNG")
    return out.getvalue()
