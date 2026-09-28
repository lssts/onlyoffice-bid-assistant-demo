"""Workflow contract checks use isolated storage, with actual DOCX generation."""
import io
import zipfile
import json
import httpx

from docx import Document
from docx.oxml.ns import qn

import app as host
from test_app import client
from workflow_documents import fill_template, inspect_controls, sample_tender


def prepared(client):
    p = client.post('/api/workflow/sample').json()
    response = client.post(f"/api/workflow/projects/{p['id']}/analyze", json={'mode':'rules'})
    assert response.status_code == 200
    p = response.json()
    response = client.post(f"/api/workflow/projects/{p['id']}/generate", json={'templates':p['analysis']['templates']})
    assert response.status_code == 200, response.text
    return response.json()


def reviewed(client):
    p = prepared(client)
    for t in p['templates']:
        vid = client.get(f"/api/documents/{t['document_id']}/versions").json()[0]['id']
        r = client.post(f"/api/workflow/projects/{p['id']}/templates/{t['id']}/review", json={
            'version_id':vid, 'mappings':{f['tag']:f['fieldKey'] for f in t['fields']}})
        assert r.status_code == 200, r.text
        p = r.json()
    return p


def test_generated_templates_preserve_inline_fields_and_table(client):
    p = prepared(client)
    assert len(p['templates']) == 3
    assert len({f['tag'] for t in p['templates'] for f in t['fields']}) == 19
    for t in p['templates']:
        data = client.get(f"/api/documents/{t['document_id']}/download").content
        doc = Document(io.BytesIO(data))
        assert len(inspect_controls(data)) == len(t['fields'])
        assert all(sdt.getparent().tag == qn('w:p') for sdt in doc.element.body.iter(qn('w:sdt')))
        if '承诺书' in t['title']:
            assert any(len(p.findall(qn('w:sdt'))) == 3 for p in doc.element.body.iter(qn('w:p')))
            assert any(r.find(qn('w:u')) is not None for r in doc.element.body.iter(qn('w:rPr')))
            assert not doc.tables
            assert '授权范围' not in ''.join(doc.element.body.itertext())
        if '财务状况' in t['title']:
            assert len(doc.tables) == 1
            assert len(list(doc.tables[0]._tbl.iter(qn('w:sdt')))) == 6
            assert len(list(doc.tables[0]._tbl.iter(qn('w:gridSpan')))) == 1


def test_complete_workflow_uses_frozen_templates_and_checked_versions(client):
    p = reviewed(client)
    original = {t['id']:client.get(f"/api/documents/{t['document_id']}/download").content for t in p['templates']}
    r = client.post(f"/api/workflow/projects/{p['id']}/fill", json={})
    assert r.status_code == 200, r.text
    p = r.json()
    assert p['fill_run']['matched'] == 19 and p['fill_run']['missing'] == 0
    for t in p['templates']:
        assert client.get(f"/api/documents/{t['document_id']}/download").content == original[t['id']]
        data = client.get(f"/api/documents/{t['output_id']}/download").content
        assert len(inspect_controls(data)) == len(t['fields'])
        assert '【' not in ''.join(v[0] for v in inspect_controls(data).values())
        vid = client.get(f"/api/documents/{t['output_id']}/versions").json()[0]['id']
        checked = client.post(f"/api/workflow/projects/{p['id']}/check", json={'template_id':t['id'],'version_id':vid})
        assert checked.status_code == 200
        assert checked.json()['passed'] and not checked.json()['issues']
    bundle = client.get(f"/api/workflow/projects/{p['id']}/bundle.zip")
    assert bundle.status_code == 200
    with zipfile.ZipFile(io.BytesIO(bundle.content)) as z:
        assert len([n for n in z.namelist() if n.endswith('.docx')]) == 3
        assert '字段与检查清单.json' in z.namelist()


def test_missing_data_is_not_invented_and_blocks_delivery(client):
    p = reviewed(client)
    client.post('/api/workflow/profiles/demo-company',json={'values':{'bidder.phone':''}})
    p = client.post(f"/api/workflow/projects/{p['id']}/fill",json={}).json()
    assert p['fill_run']['missing'] == 1
    for t in p['templates']:
        vid=client.get(f"/api/documents/{t['output_id']}/versions").json()[0]['id']
        report=client.post(f"/api/workflow/projects/{p['id']}/check",json={'template_id':t['id'],'version_id':vid}).json()
        if '授权' in t['title']:
            assert not report['passed']
            assert report['issues'][0]['message'] == '必填内容尚未补充'
    assert client.get(f"/api/workflow/projects/{p['id']}/bundle.zip").status_code == 409


def test_bad_plan_and_unreviewed_fill_are_rejected(client):
    p=client.post('/api/workflow/sample').json()
    p=client.post(f"/api/workflow/projects/{p['id']}/analyze",json={}).json()
    plan=p['analysis']['templates']
    plan[0]['fields'][0]['anchor']='invented text'
    assert client.post(f"/api/workflow/projects/{p['id']}/generate",json={'templates':plan}).status_code == 400
    assert not client.get(f"/api/workflow/projects/{p['id']}").json()['templates']
    assert client.post(f"/api/workflow/projects/{p['id']}/fill",json={}).status_code == 409


def test_removed_controls_fail_review_and_foreign_version_is_rejected(client):
    p=prepared(client)
    t=p['templates'][0]
    foreign=client.get(f"/api/documents/{p['templates'][1]['document_id']}/versions").json()[0]['id']
    body={'version_id':foreign,'mappings':{f['tag']:f['fieldKey'] for f in t['fields']}}
    url=f"/api/workflow/projects/{p['id']}/templates/{t['id']}/review"
    assert client.post(url,json=body).status_code==400
    vid=host.add_version(t['document_id'],sample_tender(),'without fields','test')
    body['version_id']=vid
    assert client.post(url,json=body).status_code==409


def test_uploaded_docx_is_actually_used_and_manual_split_supported(client):
    doc=Document()
    doc.add_paragraph('定制表单')
    doc.add_paragraph('项目名称：______；投标人名称：______')
    out=io.BytesIO();doc.save(out)
    p=client.post('/api/workflow/upload',files={'file':('custom.docx',out.getvalue())}).json()
    p=client.post(f"/api/workflow/projects/{p['id']}/analyze",json={}).json()
    assert p['analysis']['templates']==[]
    plan=client.post(f"/api/workflow/projects/{p['id']}/detect-fields",json={'templates':[{'title':'定制表单','start':0,'end':1}]}).json()['templates']
    assert {f['fieldKey'] for f in plan[0]['fields']}=={'project.name','bidder.name'}
    result=client.post(f"/api/workflow/projects/{p['id']}/generate",json={'templates':plan})
    assert result.status_code==200, result.text
    assert len(result.json()['templates'])==1


def test_pdf_and_unconfigured_ai_are_explicitly_rejected(client,monkeypatch):
    monkeypatch.delenv('BID_AI_API_KEY',raising=False)
    assert client.post('/api/workflow/upload',files={'file':('a.pdf',b'%PDF')}).status_code==400
    p=client.post('/api/workflow/sample').json()
    response=client.post(f"/api/workflow/projects/{p['id']}/analyze",json={'mode':'ai'})
    assert response.status_code==400
    assert not client.get(f"/api/workflow/projects/{p['id']}").json()['analysis']


def test_ai_adapter_validates_response_without_calling_external_service(client, monkeypatch):
    monkeypatch.setenv('BID_AI_BASE_URL', 'https://model.example/v1')
    monkeypatch.setenv('BID_AI_MODEL', 'test-model')
    monkeypatch.setenv('BID_AI_API_KEY', 'test-secret')
    doc=Document();doc.add_heading('模板',1)
    doc.add_paragraph('【投标人名称】')
    stream=io.BytesIO();doc.save(stream)
    p=client.post('/api/workflow/upload',files={'file':('ai.docx',stream.getvalue())}).json()
    payload={'templates':[{'title':'承诺书','start':0,'end':1,'fields':[
        {'block':1,'paragraph':0,'anchor':'【投标人名称】','fieldKey':'bidder.name','label':'投标人名称'}]}],
        'requirements':[]}
    class FakeAIClient:
        def __init__(self,**kwargs):pass
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        async def post(self,url,**kwargs):
            assert url=='https://model.example/v1/chat/completions'
            assert kwargs['json']['model']=='test-model'
            sent=json.loads(kwargs['json']['messages'][1]['content'])
            assert 'outline' not in sent
            result=payload
            assert [b['id'] for b in sent['blocks']]==[0,1]
            assert all('text' not in b for b in sent['blocks'])
            return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps(result,ensure_ascii=False)}}]},request=httpx.Request('POST',url))
    monkeypatch.setattr('workflow.httpx.AsyncClient',FakeAIClient)
    url=f"/api/workflow/projects/{p['id']}/analyze"
    result=client.post(url,json={'mode':'ai'})
    assert result.status_code==200, result.text
    assert result.json()['analysis']['mode']=='ai'
    payload['templates'][0]['fields'][0]['anchor']='原文不存在的字段'
    result=client.post(url,json={'mode':'ai'})
    assert result.status_code==502
    assert 'test-secret' not in result.text
    assert '投标人名称' in result.json()['detail']
    assert '锚点' in result.json()['detail']


def test_reopening_template_creates_editable_copy_and_requires_review_again(client):
    p=reviewed(client)
    t=p['templates'][0]
    r=client.post(f"/api/workflow/projects/{p['id']}/templates/{t['id']}/reopen")
    assert r.status_code==200
    updated=r.json()['templates'][0]
    assert updated['document_id']!=t['document_id'] and not updated['reviewed']
    assert client.get('/api/versions/'+t['review_version_id']+'/download').status_code==200
    assert client.post(f"/api/workflow/projects/{p['id']}/fill",json={}).status_code==409


def test_outline_ranges_and_direct_level_override():
    from docx.oxml import OxmlElement
    from workflow_documents import parse_source, extract_outline
    doc=Document()
    doc.add_heading('第一章 投标文件格式',1)
    p=doc.add_heading('财务状况表',2)
    node=OxmlElement('w:outlineLvl'); node.set(qn('w:val'),'2'); p._p.get_or_add_pPr().append(node)
    doc.add_table(rows=1, cols=1).cell(0,0).text='待填写'
    doc.add_heading('第二章 技术规范',1)
    stream=io.BytesIO(); doc.save(stream)
    outline=extract_outline(parse_source(stream.getvalue()))
    assert [(c['start'],c['end'],c['level']) for c in outline]==[(0,2,1),(1,2,3),(3,3,1)]
    assert outline[1]['parent_id']==outline[0]['id']


def test_ai_without_outline_does_not_call_model(monkeypatch):
    import asyncio
    import workflow_ai
    async def fake(*args):
        raise AssertionError("should not call model")
    monkeypatch.setattr(workflow_ai,'request_model',fake)
    result=asyncio.run(workflow_ai.analyze_chapters({'blocks':[], 'outline':[]},lambda s,t:t))
    assert result['scope']['model_calls']==0 and result['templates']==[]


def test_generate_more_than_thirty_templates(client):
    doc=Document()
    for i in range(31):doc.add_paragraph(f'模板 {i+1}')
    stream=io.BytesIO();doc.save(stream)
    response=client.post('/api/workflow/upload',files={'file':('many.docx',stream.getvalue())})
    assert response.status_code==200,response.text
    p=response.json()
    assert client.post(f"/api/workflow/projects/{p['id']}/analyze",json={'mode':'rules'}).status_code==200
    plan=[{'title':f'模板{i+1}','start':i,'end':i,'fields':[]} for i in range(31)]
    response=client.post(f"/api/workflow/projects/{p['id']}/generate",json={'templates':plan})
    assert response.status_code==200,response.text
    assert len(response.json()['templates'])==31


def test_mark_span_preserves_surrounding_objects_and_multiple_fields():
    from docx.oxml import OxmlElement
    from workflow_documents import mark_span, text_of
    from lxml import etree
    doc=Document(); p=doc.add_paragraph()
    r=p.add_run('名称：'); r.add_break(); r.add_text('【名称】'); r.add_tab(); r.add_text('日期：【日期】')
    drawing=OxmlElement('w:drawing');r._r.append(drawing)
    bookmark=OxmlElement('w:bookmarkStart');bookmark.set(qn('w:id'),'8');bookmark.set(qn('w:name'),'keep');p._p.append(bookmark)
    fr=p.add_run()._r
    for kind in ('begin','separate'):
        node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),kind);fr.append(node)
    t=OxmlElement('w:t');t.text='1';fr.append(t)
    node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),'end');fr.append(node)
    original={tag:[etree.tostring(n) for n in p._p.iter(qn('w:'+tag))] for tag in ('br','tab','drawing','bookmarkStart','fldChar')}
    text=text_of(p._p)
    for label in ('日期','名称'):
        anchor='【'+label+'】'; start=text.index(anchor)
        mark_span(p._p,start,start+len(anchor),'field:'+label,label,123)
    assert text_of(p._p)==text
    assert len(list(p._p.iter(qn('w:sdt'))))==2
    for tag, nodes in original.items():
        assert [etree.tostring(n) for n in p._p.iter(qn('w:'+tag))]==nodes


def test_mark_span_rejects_crossing_break_or_field_without_mutation():
    import pytest
    from docx.oxml import OxmlElement
    from workflow_documents import mark_span
    doc=Document();p=doc.add_paragraph();r=p.add_run('AA');r.add_break();r.add_text('BB')
    before=p._p.xml
    with pytest.raises(ValueError,match='跨越'):
        mark_span(p._p,0,4,'field:x','x',1)
    assert p._p.xml==before
    doc=Document();p=doc.add_paragraph();r=p.add_run()._r
    node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),'begin');r.append(node)
    t=OxmlElement('w:t');t.text='AA';r.append(t)
    before=p._p.xml
    with pytest.raises(ValueError,match='复杂域'):
        mark_span(p._p,0,2,'field:x','x',1)
    assert p._p.xml==before


def test_ai_direct_last_parent_includes_children_without_title_screening(monkeypatch):
    import asyncio
    import workflow_ai
    from workflow_documents import parse_source, extract_outline
    doc=Document();doc.add_heading('工程概况',1);doc.add_paragraph('无关正文')
    doc.add_heading('投标文件格式',1);doc.add_heading('承诺书',2);doc.add_paragraph('模板正文')
    stream=io.BytesIO();doc.save(stream);blocks=parse_source(stream.getvalue())
    calls=[]
    async def fake(instruction,payload):
        calls.append(payload)
        assert 'outline' not in payload
        assert [b['id'] for b in payload['blocks']]==[2,3,4]
        return {'templates':[]}
    monkeypatch.setattr(workflow_ai,'request_model',fake)
    result=asyncio.run(workflow_ai.analyze_chapters({'blocks':blocks,'outline':extract_outline(blocks)},lambda s,t:t))
    assert len(calls)==1 and result['scope']['model_calls']==1
    assert result['scope']['strategy']=='last_parent_direct'


def test_ai_empty_cell_is_generated_and_fillable(client, monkeypatch):
    import workflow_ai
    from workflow_documents import inspect_controls, fill_template
    monkeypatch.setenv('BID_AI_BASE_URL','https://model.example/v1')
    monkeypatch.setenv('BID_AI_MODEL','test');monkeypatch.setenv('BID_AI_API_KEY','test')
    doc=Document();doc.add_heading('最后一章名称无需匹配',1)
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='投标人名称'
    stream=io.BytesIO();doc.save(stream)
    p=client.post('/api/workflow/upload',files={'file':('empty.docx',stream.getvalue())}).json()
    async def fake(instruction,payload):
        paragraph=payload['blocks'][1]['paragraphs'][1]
        assert paragraph['empty_cell'] and paragraph['cell']['column']==1
        return {'templates':[{'title':'企业信息','start':0,'end':1,'fields':[
            {'position_id':paragraph['positions'][0]['id'],'label':'投标人名称','fieldKey':'bidder.name'}]}]}
    monkeypatch.setattr(workflow_ai,'request_model',fake)
    response=client.post(f"/api/workflow/projects/{p['id']}/analyze",json={'mode':'ai'})
    assert response.status_code==200,response.text
    plan=response.json()['analysis']['templates']
    response=client.post(f"/api/workflow/projects/{p['id']}/generate",json={'templates':plan})
    assert response.status_code==200,response.text
    t=response.json()['templates'][0];tag=t['fields'][0]['tag']
    data=client.get(f"/api/documents/{t['document_id']}/download").content
    assert inspect_controls(data)[tag]==['【投标人名称】']
    filled=fill_template(data,{tag:'测试企业'})
    assert inspect_controls(filled)[tag]==['测试企业']
    assert Document(io.BytesIO(filled)).tables[0].cell(0,0).text=='投标人名称'


def test_empty_cell_duplicate_and_nonempty_cell_rejected(client):
    doc=Document();doc.add_heading('模板',1)
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='字段标签'
    stream=io.BytesIO();doc.save(stream)
    p=client.post('/api/workflow/upload',files={'file':('empty.docx',stream.getvalue())}).json()
    client.post(f"/api/workflow/projects/{p['id']}/analyze",json={'mode':'rules'})
    field={'block':1,'paragraph':1,'anchor':'','kind':'empty_cell','start':0,'end':0,'label':'测试'}
    plan=[{'title':'测试','start':0,'end':1,'fields':[field,field.copy()]}]
    url=f"/api/workflow/projects/{p['id']}/generate"
    assert client.post(url,json={'templates':plan}).status_code==400
    plan[0]['fields']=[{**field,'paragraph':0}]
    assert client.post(url,json={'templates':plan}).status_code==400


def test_whitespace_anchor_matches_whole_runs_only():
    import pytest
    from workflow_ai import anchor_match
    text='日期：    年     月  日'
    assert anchor_match(text,'    ',0).span()==(3,7)
    assert anchor_match(text,'     ',0).span()==(8,13)
    assert anchor_match(text,'  ',0).span()==(14,16)
    with pytest.raises(ValueError,match='位置 ID'):
        anchor_match(text,'   ',0)
    assert anchor_match('姓名：    联系人：    ','    ',1).start()==11


def test_position_id_resolves_exact_blank_and_rejects_unknown_id(monkeypatch):
    import asyncio
    import pytest
    import workflow_ai
    from fastapi import HTTPException
    from workflow_documents import parse_source,extract_outline
    doc=Document();doc.add_heading('模板',1);doc.add_paragraph('日期：    年     月  日')
    stream=io.BytesIO();doc.save(stream);blocks=parse_source(stream.getvalue())
    state={'blocks':blocks,'outline':extract_outline(blocks)}
    invalid=False
    async def fake(instruction,payload):
        positions=payload['blocks'][1]['paragraphs'][0]['positions']
        return {'templates':[{'title':'模板','start':0,'end':1,'fields':[
            {'position_id':'invalid' if invalid else positions[2]['id'],'label':'日期'}]}]}
    monkeypatch.setattr(workflow_ai,'request_model',fake)
    result=asyncio.run(workflow_ai.analyze_chapters(state,lambda s,t:t))
    field=result['templates'][0]['fields'][0]
    assert (field['start'],field['end'],field['anchor'])==(14,16,'  ')
    invalid=True
    with pytest.raises(HTTPException) as exc:
        asyncio.run(workflow_ai.analyze_chapters(state,lambda s,t:t))
    assert '未知的 position_id' in exc.value.detail


def test_complete_placeholder_replaced_and_seal_instructions_kept():
    from workflow_documents import generate_template, fill_template, text_of
    doc=Document();p=doc.add_paragraph();p.add_run('致：    （招标');p.add_run('人名称）')
    doc.add_paragraph('投标人：    （盖单位章）')
    stream=io.BytesIO();doc.save(stream)
    fields=[{'id':'12345678','tag':'field:owner','label':'招标人名称','block':0,'paragraph':0,'start':2,'end':6},
            {'id':'23456789','tag':'field:bidder','label':'投标人名称','fieldKey':'bidder.name','block':1,'paragraph':0,'start':4,'end':8}]
    data=generate_template(stream.getvalue(),{'start':0,'end':1,'fields':fields})
    result=fill_template(data,{'field:owner':'测试招标单位','field:bidder':'测试投标单位'})
    paragraphs=Document(io.BytesIO(result)).paragraphs
    assert text_of(paragraphs[0]._p)=='致：测试招标单位'
    assert text_of(paragraphs[1]._p)=='投标人：测试投标单位（盖单位章）'


def test_date_placeholder_is_short_and_negative_spacing_removed():
    from docx.oxml import OxmlElement
    from workflow_documents import generate_template,fill_template,text_of
    doc=Document();p=doc.add_paragraph('日期：')
    for blank, unit in [('    ','年'),('   ','月'),('  ','日')]:
        run=p.add_run(blank);run.bold=True
        spacing=OxmlElement('w:spacing');spacing.set(qn('w:val'),'-98');run._r.get_or_add_rPr().append(spacing)
        p.add_run(unit)
    stream=io.BytesIO();doc.save(stream)
    fields=[{'id':f'{i+1:08x}','tag':f'field:{i}','label':label,'block':0,'paragraph':0,'start':a,'end':b}
            for i,(a,b,label) in enumerate([(3,7,'签署日期年份'),(8,11,'签署日期月份'),(12,14,'签署日期日')])]
    data=generate_template(stream.getvalue(),{'start':0,'end':0,'fields':fields})
    assert text_of(Document(io.BytesIO(data)).paragraphs[0]._p)=='日期：____年__月__日'
    result=fill_template(data,{'field:0':'2026','field:1':'09','field:2':'28'})
    p=Document(io.BytesIO(result)).paragraphs[0]._p
    assert text_of(p)=='日期：2026年09月28日'
    assert not list(p.iter(qn('w:spacing')))
    assert len(list(p.iter(qn('w:b'))))==3


def test_repair_existing_copy_preserves_user_edits_and_filled_value():
    from workflow_documents import mark_span,repair_template_placeholders,text_of
    doc=Document();p=doc.add_paragraph('用户修改的致辞：    （招标人名称）')
    start=p.text.index('    ');mark_span(p._p,start,start+4,'field:owner','招标人名称',123,placeholder='用户填写公司')
    stream=io.BytesIO();doc.save(stream)
    result=repair_template_placeholders(stream.getvalue(),[{'tag':'field:owner','label':'招标人名称'}])
    assert text_of(Document(io.BytesIO(result)).paragraphs[0]._p)=='用户修改的致辞：用户填写公司'


def test_repair_endpoint_creates_copy_and_requires_review_again(client):
    p=reviewed(client);t=p['templates'][0];original=t['document_id']
    r=client.post(f"/api/workflow/projects/{p['id']}/templates/{t['id']}/repair-placeholders",json={
        'version_id':t['review_version_id'],'mappings':{f['tag']:f['fieldKey'] for f in t['fields']}})
    assert r.status_code==200,r.text
    repaired=r.json()['templates'][0]
    assert repaired['document_id']!=original and not repaired['reviewed']
    assert client.get('/api/versions/'+t['review_version_id']+'/download').status_code==200
