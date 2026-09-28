"""Source-preserving DOCX splitting, inline SDTs and deterministic demo analysis."""
from __future__ import annotations

import io
import re
import uuid
from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def heading_info(doc, element, styles=None):
    """Read actual heading metadata; plain-title heuristics never inspect body text."""
    if element.tag != qn("w:p"):
        return None
    text = text_of(element).strip()
    if not text or len(text) > 200:
        return None
    ppr = element.find(qn("w:pPr"))
    style_node = ppr.find(qn("w:pStyle")) if ppr is not None else None
    style_id = style_node.get(qn("w:val")) if style_node is not None else ""
    styles = styles if styles is not None else {s.style_id: s for s in doc.styles}
    style = styles.get(style_id)
    style_name = style.name if style is not None else style_id
    if re.search(r"^(?:toc|目录)", style_name, re.I) or re.search(r"(?:\.{3,}|…{2,}|\t)\s*\d+\s*$", text):
        return None
    if list(element.iter(qn("w:fldChar"))) or list(element.iter(qn("w:instrText"))):
        return None
    props, seen = [(ppr, "paragraph-outline")], set()
    fallback_level = None
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        props.append((style.element.find(qn("w:pPr")), "heading-style"))
        match = re.match(r"^(?:Heading\s*|标题\s*)([1-9])$", style.name, re.I)
        if match and fallback_level is None:
            fallback_level = int(match.group(1))
        style = style.base_style
    for pr, origin in props:
        node = pr.find(qn("w:outlineLvl")) if pr is not None else None
        if node is not None:
            value = int(node.get(qn("w:val"), "9"))
            if 0 <= value <= 8:
                return {"level": value + 1, "source": origin}
            return None
    if fallback_level is not None:
        return {"level": fallback_level, "source": "heading-style"}
    if len(text) > 90 or re.search(r"[。；;]", text):
        return None
    if re.match(r"^第[一二三四五六七八九十百零〇\d]+[编篇章]\s*\S", text):
        return {"level": 1, "source": "title-pattern"}
    if re.match(r"^第[一二三四五六七八九十百零〇\d]+节\s*\S", text):
        return {"level": 2, "source": "title-pattern"}
    if re.match(r"^(?:附件|附录|格式|模板)\s*[一二三四五六七八九十\d]+\s*[：:、.．\s]", text):
        return {"level": 2, "source": "title-pattern"}
    return None


def extract_outline(blocks):
    """A heading's range ends immediately before the next peer/ancestor heading."""
    outline, stack = [], []
    prefix = [0]
    for block in blocks:
        prefix.append(prefix[-1] + len(block["text"]))
    for block in blocks:
        heading = block.get("heading")
        if not heading:
            continue
        while stack and stack[-1]["level"] >= heading["level"]:
            stack.pop()["end"] = block["id"] - 1
        item = {"id": f"chapter-{block['id']}", "title": block["text"].strip(),
                "start": block["id"], "end": len(blocks) - 1, **heading,
                "parent_id": stack[-1]["id"] if stack else None}
        outline.append(item)
        stack.append(item)
    for item in outline:
        item["characters"] = prefix[item["end"] + 1] - prefix[item["start"]]
    return outline


CATALOG = [
    {"key": "bidder.name", "label": "投标人名称", "aliases": ["投标人名称", "公司名称", "企业名称", "投标人"], "source": "企业档案 / 企业名称"},
    {"key": "bidder.legal_representative", "label": "法定代表人", "aliases": ["法定代表人", "法人姓名"], "source": "企业档案 / 法定代表人"},
    {"key": "bidder.address", "label": "企业地址", "aliases": ["企业地址", "注册地址", "地址"], "source": "企业档案 / 注册地址"},
    {"key": "bidder.phone", "label": "联系电话", "aliases": ["联系电话", "电话"], "source": "企业档案 / 联系电话"},
    {"key": "project.name", "label": "项目名称", "aliases": ["项目名称", "工程名称"], "source": "项目档案 / 项目名称"},
    {"key": "project.number", "label": "项目编号", "aliases": ["项目编号", "招标编号"], "source": "项目档案 / 招标编号"},
    {"key": "agent.name", "label": "代理人姓名", "aliases": ["代理人姓名", "被授权人", "委托代理人"], "source": "人员档案 / 授权代理人"},
    {"key": "agent.id_number", "label": "代理人身份证号", "aliases": ["代理人身份证号", "身份证号码", "身份证号"], "source": "人员档案 / 身份证号"},
    {"key": "sign.date", "label": "签署日期", "aliases": ["签署日期", "日期"], "source": "项目档案 / 签署日期"},
] + [{"key": f"finance.{year}.{metric}", "label": f"{year}年{label}", "aliases": [f"{year}年{label}"],
      "source": f"财务档案 / {year} / {label}（万元）"}
     for year in (2023, 2024, 2025) for metric, label in (("revenue", "营业收入"), ("profit", "净利润"), ("assets", "资产总额"))]
CATALOG_BY_KEY = {f["key"]: f for f in CATALOG}
ALIASES = {a: f["key"] for f in CATALOG for a in f["aliases"]}
PLACEHOLDER = re.compile(r"【([^【】\n]{1,40})】|\{\{([^{}\n]{1,60})\}\}|[_＿]{2,}")


def text_of(element):
    return "".join(n.text or "" for n in element.iter(qn("w:t")))


def body_blocks(doc):
    return [n for n in doc.element.body if n.tag != qn("w:sectPr")]


def table_paragraph_metadata(table):
    """Expose physical cell coordinates without duplicating merged cell text."""
    metadata = {}
    for ti, tbl in enumerate(table.iter(qn("w:tbl"))):
        for ri, row in enumerate(tbl.findall(qn("w:tr"))):
            before = row.find("./" + qn("w:trPr") + "/" + qn("w:gridBefore"))
            column = int(before.get(qn("w:val"))) if before is not None else 0
            for cell in row.findall(qn("w:tc")):
                span = cell.find("./" + qn("w:tcPr") + "/" + qn("w:gridSpan"))
                width = int(span.get(qn("w:val"))) if span is not None else 1
                merge = cell.find("./" + qn("w:tcPr") + "/" + qn("w:vMerge"))
                vertical = merge.get(qn("w:val"), "continue") if merge is not None else None
                paragraphs = cell.findall(qn("w:p"))
                safe = not text_of(cell).strip() and vertical != "continue" and all(
                    n.tag in {qn("w:pPr"), qn("w:r")} for p in paragraphs for n in p) and all(
                    n.tag in {qn("w:rPr"), qn("w:t")} for p in paragraphs for r in p.findall(qn("w:r")) for n in r)
                safe = safe and all(n.tag in {qn("w:tcPr"), qn("w:p")} for n in cell)
                for pi, paragraph in enumerate(paragraphs):
                    metadata[paragraph] = {"cell": {"table": ti, "row": ri, "column": column,
                        "grid_span": width, "vertical_merge": vertical}, "empty_cell": bool(safe and pi == 0)}
                column += width
    return metadata


def parse_source(data):
    doc = Document(io.BytesIO(data))
    blocks = []
    styles = {s.style_id: s for s in doc.styles}
    for i, el in enumerate(body_blocks(doc)):
        paragraphs = [el] if el.tag == qn("w:p") else list(el.iter(qn("w:p")))
        metadata = table_paragraph_metadata(el) if el.tag == qn("w:tbl") else {}
        blocks.append({"id": i, "kind": "table" if el.tag == qn("w:tbl") else "paragraph",
                       "text": "\n".join(text_of(p) for p in paragraphs),
                       "heading": heading_info(doc, el, styles),
                       "paragraphs": [{"index": j, "text": text_of(p), **metadata.get(p, {})} for j, p in enumerate(paragraphs)]})
    return blocks


def guess_fields(blocks, start, end):
    result = []
    for b in blocks[start:end + 1]:
        for p in b["paragraphs"]:
            for match in PLACEHOLDER.finditer(p["text"]):
                label = match.group(1) or match.group(2)
                if not label:
                    prefix = p["text"][:match.start()].rstrip(" ：:\t")
                    label = next((a for a in sorted(ALIASES, key=len, reverse=True) if prefix.endswith(a)), "未识别字段")
                key = ALIASES.get(label, label if label in CATALOG_BY_KEY else "")
                result.append({"id": uuid.uuid4().hex[:12], "block": b["id"], "paragraph": p["index"],
                               "start": match.start(), "end": match.end(), "anchor": match.group(),
                               "label": label, "fieldKey": key, "required": True})
    return result


def rule_analysis(blocks):
    starts = []
    for b in blocks:
        t = b["text"].strip()
        if b["kind"] == "paragraph" and len(t) <= 65 and (
            re.match(r"^(?:附件|格式|模板)\s*[一二三四五六七八九十\d]+\s*[：:、.．\s]", t)
            or re.match(r"^(?:[一二三四五六七八九十\d]+[、.．]\s*)?(?:投标承诺书|承诺书|财务状况表|法定代表人授权委托书|授权委托书)$", t)
        ):
            starts.append(b["id"])
    templates = []
    for i, start in enumerate(starts):
        end = starts[i + 1] - 1 if i + 1 < len(starts) else len(blocks) - 1
        templates.append({"id": uuid.uuid4().hex[:12], "title": blocks[start]["text"].strip(), "start": start, "end": end,
                          "reason": f"根据原文第 {start + 1} 块标题识别，范围需人工确认。",
                          "fields": guess_fields(blocks, start, end)})
    requirements = [{"block": b["id"], "quote": b["text"][:1000]} for b in blocks
                    if re.search(r"字体|字号|页边距|行距|格式要求|不得修改|不得更改|万元", b["text"])][:30]
    return {"templates": templates, "requirements": requirements,
            "warnings": ["本地规则只识别常见模板标题，以及【字段】、{{字段}}和下划线占位；空白单元格、复杂语义与扫描件不在自动识别范围。",
                         "拆分保留来源样式、表格与资源；分页、分节、浮动对象及原文格式是否符合招标要求仍需人工核对。"]}


def run_with_text(text, rpr=None):
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(deepcopy(rpr))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def field_rpr(rpr):
    """Blank-specific compression must not squeeze the replacement text."""
    if rpr is None:
        return None
    result = deepcopy(rpr)
    for node in list(result):
        if (node.tag == qn("w:spacing") and int(node.get(qn("w:val"), "0")) < 0) or node.tag == qn("w:fitText"):
            result.remove(node)
    return result


def matching_hint(hint, label, key=""):
    # Only matching field names are placeholders; signature/seal instructions stay.
    if re.search(r"签字|签章|盖章|盖.*章|复印件|扫描件|如有|适用|注[:：]", hint):
        return False
    normalize = lambda value: re.sub(r"[\s（）()【】]", "", value)
    names = [label] + CATALOG_BY_KEY.get(key, {}).get("aliases", [])
    return normalize(hint) in {normalize(name) for name in names if name}


def replacement_span(text, start, end, label, key=""):
    # Expand a blank to include its semantic hint, including the parentheses.
    suffix = re.match(r"[ \u3000\u00a0]*[（(]([^（）()\r\n]{1,60})[）)]", text[end:])
    if suffix and matching_hint(suffix.group(1), label, key):
        end += suffix.end()
    # Also handle a model selecting just the name inside parentheses.
    if start > 0 and end < len(text) and text[start-1] in "（(" and text[end] in "）)" and matching_hint(text[start:end], label, key):
        start -= 1
        end += 1
    return start, end


def field_placeholder(label, following):
    unit = re.match(r"^[ \u3000\u00a0]*([年月日])", following)
    if unit:
        return "____" if unit.group(1) == "年" else "__"
    return "【" + label + "】"


def remove_plain_text(paragraph, start, end):
    """Remove hint text across runs, preserving every surrounding OOXML object."""
    offset, edits = 0, []
    for child in paragraph:
        length = len(text_of(child))
        if child.tag != qn("w:r"):
            if offset < end and offset + length > start:
                raise ValueError("占位提示位于复杂对象中，无法安全删除")
            offset += length
            continue
        for node in child:
            if node.tag == qn("w:rPr"):
                continue
            value = node.text or "" if node.tag == qn("w:t") else text_of(node)
            finish = offset + len(value)
            if offset < end and finish > start:
                if node.tag != qn("w:t"):
                    raise ValueError("占位提示跨越复杂对象，无法安全删除")
                edits.append((node, value[:max(0,start-offset)] + value[min(len(value),end-offset):]))
            elif node.tag != qn("w:t") and start < offset < end:
                raise ValueError("占位提示跨越图片或换行，无法安全删除")
            offset = finish
    for node, value in edits:
        node.text = value


def repair_template_placeholders(data, fields):
    """Repair a saved template copy without regenerating or erasing user edits."""
    doc = Document(io.BytesIO(data))
    lookup = {f["tag"]: f for f in fields}
    for control in doc.element.body.iter(qn("w:sdt")):
        tag = control.find("./" + qn("w:sdtPr") + "/" + qn("w:tag"))
        field = lookup.get(tag.get(qn("w:val"))) if tag is not None else None
        parent = control.getparent()
        if not field or parent.tag != qn("w:p"):
            continue
        content = control.find(qn("w:sdtContent"))
        if content is None:
            continue
        before = sum(len(text_of(n)) for n in parent[:parent.index(control)])
        end = before + len(text_of(control))
        original = text_of(parent)
        _, expanded = replacement_span(original, end, end, field["label"], field.get("fieldKey", ""))
        if expanded > end:
            remove_plain_text(parent, end, expanded)
        if text_of(content) == "【" + field["label"] + "】":
            placeholder = field_placeholder(field["label"], text_of(parent)[end:])
            if placeholder != text_of(content):
                rpr = next(content.iter(qn("w:rPr")), None)
                replacement = run_with_text(placeholder, field_rpr(rpr))
                for node in list(content):
                    content.remove(node)
                content.append(replacement)
        for rpr in list(content.iter(qn("w:rPr"))):
            rpr.getparent().replace(rpr, field_rpr(rpr))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def mark_span(paragraph, start, end, tag, label, control_id, placeholder=None):
    """Replace only a plain text span; retain surrounding OOXML objects in order."""
    atoms, offset, field_depth = [], 0, 0
    for child in paragraph:
        if child.tag != qn("w:r"):
            length = len(text_of(child))
            atoms.append((deepcopy(child), offset, offset + length, False))
            offset += length
            continue
        rpr = child.find(qn("w:rPr"))
        for node in child:
            if node.tag == qn("w:rPr"):
                continue
            if node.tag == qn("w:fldChar") and node.get(qn("w:fldCharType")) == "begin":
                field_depth += 1
            run = OxmlElement("w:r")
            run.attrib.update(child.attrib)
            if rpr is not None:
                run.append(deepcopy(rpr))
            run.append(deepcopy(node))
            length = len(node.text or "") if node.tag == qn("w:t") else len(text_of(node))
            atoms.append((run, offset, offset + length, node.tag == qn("w:t") and field_depth == 0))
            offset += length
            if node.tag == qn("w:fldChar") and node.get(qn("w:fldCharType")) == "end":
                field_depth = max(0, field_depth - 1)
    if not 0 <= start < end <= offset:
        raise ValueError("字段文字范围无效")
    for node, a, b, plain in atoms:
        if not plain and ((a < end and b > start) or (a == b and start < a < end)):
            raise ValueError(f"字段“{label}”跨越图片、换行、书签或位于复杂域/超链接/已有控件中，请调整该字段的文字范围")
    first = next((n for n, a, b, plain in atoms if plain and a < end and b > start), None)
    if first is None:
        raise ValueError("没有找到字段对应的普通文字")
    sdt = OxmlElement("w:sdt")
    props = OxmlElement("w:sdtPr")
    for name, value in (("tag", tag), ("alias", label), ("id", str(control_id))):
        n = OxmlElement("w:" + name)
        n.set(qn("w:val"), value)
        props.append(n)
    content = OxmlElement("w:sdtContent")
    content.append(run_with_text(placeholder if placeholder is not None else "【" + label + "】", field_rpr(first.find(qn("w:rPr")))))
    sdt.extend([props, content])
    result, inserted = [], False
    for node, a, b, plain in atoms:
        if plain and a < end and b > start:
            text = text_of(node)
            if a < start:
                result.append(run_with_text(text[:start-a], node.find(qn("w:rPr"))))
            if not inserted:
                result.append(sdt)
                inserted = True
            if b > end:
                result.append(run_with_text(text[end-a:], node.find(qn("w:rPr"))))
        else:
            result.append(node)
    for child in list(paragraph):
        paragraph.remove(child)
    paragraph.extend(result)


def generate_template(data, spec):
    doc = Document(io.BytesIO(data))
    blocks = body_blocks(doc)
    # Rebuild marked paragraphs in one pass using offsets in the original text.
    grouped = {}
    for f in spec["fields"]:
        grouped.setdefault((f["block"], f["paragraph"]), []).append(f)
    for (bi, pi), fields in grouped.items():
        el = blocks[bi]
        p = el if el.tag == qn("w:p") else list(el.iter(qn("w:p")))[pi]
        original = text_of(p)
        normalized = []
        for f in fields:
            a, b = f["start"], f["end"]
            if f.get("kind") != "empty_cell":
                a, b = replacement_span(original, a, b, f["label"], f.get("fieldKey", ""))
            normalized.append({**f, "start": a, "end": b})
        ordered = sorted(normalized, key=lambda f: f["start"])
        if any(a["end"] > b["start"] for a, b in zip(ordered, ordered[1:])):
            raise ValueError("完整占位区域重叠，请合并重复字段后重新生成")
        # Right-to-left insertion keeps original offsets valid for remaining fields.
        for f in reversed(ordered):
            if f.get("kind") == "empty_cell":
                # Recheck against the original OOXML before inserting into a blank cell.
                if not table_paragraph_metadata(el).get(p, {}).get("empty_cell"):
                    raise ValueError("空白单元格结构已变化，请重新识别字段")
                rpr = next(p.iter(qn("w:rPr")), None)
                if rpr is None:
                    rpr = p.find("./" + qn("w:pPr") + "/" + qn("w:rPr"))
                placeholder = run_with_text("_", rpr)
                for r in list(p.findall(qn("w:r"))):
                    p.remove(r)
                p.append(placeholder)
                mark_span(p, 0, 1, f["tag"], f["label"], int(f["id"][:7], 16))
            else:
                mark_span(p, f["start"], f["end"], f["tag"], f["label"], int(f["id"][:7], 16),
                          field_placeholder(f["label"], original[f["end"]:]))
    for i, el in enumerate(blocks):
        if not spec["start"] <= i <= spec["end"] and el.getparent() is not None:
            el.getparent().remove(el)
    first_p = next(doc.element.body.iter(qn("w:p")), None)
    if first_p is not None:
        for n in list(first_p.iter(qn("w:pageBreakBefore"))):
            n.getparent().remove(n)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def inspect_controls(data):
    doc = Document(io.BytesIO(data))
    found = {}
    for sdt in doc.element.body.iter(qn("w:sdt")):
        tag_node = sdt.find("./" + qn("w:sdtPr") + "/" + qn("w:tag"))
        if tag_node is not None:
            tag = tag_node.get(qn("w:val"))
            found.setdefault(tag, []).append(text_of(sdt.find(qn("w:sdtContent"))))
    return found


def fill_template(data, values):
    doc = Document(io.BytesIO(data))
    for sdt in doc.element.body.iter(qn("w:sdt")):
        tag_node = sdt.find("./" + qn("w:sdtPr") + "/" + qn("w:tag"))
        if tag_node is None or tag_node.get(qn("w:val")) not in values:
            continue
        value = values[tag_node.get(qn("w:val"))]
        if value is None or value == "":
            continue
        content = sdt.find(qn("w:sdtContent"))
        if any(c.tag != qn("w:r") for c in content):
            raise ValueError("回填区域已被改成复杂内容，请恢复为行内文字控件")
        rpr = next(content.iter(qn("w:rPr")), None)
        replacement = run_with_text(str(value), field_rpr(rpr))
        for child in list(content):
            content.remove(child)
        content.append(replacement)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def sample_tender():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.5)
    for name in ("Normal", "Title", "Heading 1"):
        style = doc.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.5
    doc.add_heading("示范工程招标文件", 0)
    doc.add_paragraph("本文件为业务流程演示材料，项目、企业、人员和财务数据均为虚构。")
    doc.add_paragraph("格式要求：使用 A4 纵向页面，左右页边距 2.5 厘米，正文 11 磅，1.5 倍行距。以下三份附件应分别提交。固定承诺文字不得修改，财务金额单位为万元。")
    doc.add_heading("附件一：投标承诺书", 1)
    doc.add_paragraph("投标人名称：【投标人名称】；项目名称：【项目名称】；项目编号：【项目编号】。")
    doc.add_paragraph("我公司承诺所提交的投标资料真实、完整，接受招标文件规定的评审程序，并承担相应责任。")
    p = doc.add_paragraph("法定代表人：")
    p.add_run("【法定代表人】").underline = True
    p.add_run("；签署日期：【签署日期】。")
    doc.add_heading("附件二：财务状况表", 1)
    doc.add_paragraph("投标人名称：【投标人名称】。金额单位：万元。")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for c, text in zip(table.rows[0].cells, ("财务指标", "2023 年", "2024 年")):
        c.text = text
    for label in ("营业收入", "净利润", "资产总额"):
        for c, text in zip(table.add_row().cells, (label, f"【2023年{label}】", f"【2024年{label}】")):
            c.text = text
    cell = table.add_row().cells[0]
    cell.merge(table.rows[-1].cells[2]).text = "以上数据应与相应年度审计报告一致。"
    doc.add_heading("附件三：法定代表人授权委托书", 1)
    doc.add_paragraph("本人【法定代表人】系【投标人名称】的法定代表人，现委托【代理人姓名】为我方代理人，参加【项目名称】投标。")
    doc.add_paragraph("代理人身份证号：【代理人身份证号】；联系电话：【联系电话】。")
    doc.add_paragraph("授权范围：办理本项目投标相关事宜。代理人无转委托权。")
    doc.add_paragraph("签署日期：【签署日期】。")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
