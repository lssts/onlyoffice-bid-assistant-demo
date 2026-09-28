from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
import secrets
import sqlite3
import time
import uuid
import zipfile
from contextlib import asynccontextmanager, contextmanager
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

import httpx
import jwt
from docx import Document
from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field

from samples import certificate_png, sample_docx

ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")
DATA = Path(os.getenv("DEMO_DATA_DIR", str(ROOT / "data")))
DS = os.getenv("ONLYOFFICE_URL", "http://localhost:9898").rstrip("/")
CONTAINER_BASE = os.getenv("BACKEND_CONTAINER_URL", "http://host.docker.internal:8010").rstrip("/")
APP_SECRET = os.getenv("APP_SIGNING_SECRET", "")
JWT_ENABLED = os.getenv("ONLYOFFICE_JWT_ENABLED", "true").lower() == "true"
JWT_SECRET = os.getenv("ONLYOFFICE_JWT_SECRET", "")
COMMAND_ENABLED = os.getenv("ONLYOFFICE_COMMAND_JWT_ENABLED", str(JWT_ENABLED)).lower() == "true"
COMMAND_SECRET = os.getenv("ONLYOFFICE_COMMAND_SECRET", JWT_SECRET)
OUTBOX_ENABLED = os.getenv("ONLYOFFICE_OUTBOX_JWT_ENABLED", str(JWT_ENABLED)).lower() == "true"
OUTBOX_SECRET = os.getenv("ONLYOFFICE_OUTBOX_SECRET", JWT_SECRET)
JWT_HEADER = os.getenv("ONLYOFFICE_JWT_HEADER", "Authorization")
locks: dict[str, asyncio.Lock] = {}


def now():
    return datetime.now(timezone.utc).isoformat()


@contextmanager
def db():
    con = sqlite3.connect(DATA / "demo.sqlite3", timeout=15)
    con.row_factory = sqlite3.Row
    try:
        with con:
            yield con
    finally:
        con.close()


def init_data():
    global APP_SECRET
    DATA.mkdir(parents=True, exist_ok=True)
    for folder in ["documents", "versions", "exports", "assets"]:
        (DATA / folder).mkdir(exist_ok=True)
    if not APP_SECRET:
        secret_path = DATA / ".signing-secret"
        if not secret_path.exists():
            secret_path.write_text(secrets.token_urlsafe(48), encoding="utf-8")
        APP_SECRET = secret_path.read_text(encoding="utf-8")
    with db() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS documents (
          id TEXT PRIMARY KEY, title TEXT NOT NULL, doc_key TEXT NOT NULL,
          created_at TEXT NOT NULL, saved_at TEXT, save_error TEXT, last_hash TEXT,
          active_users INTEGER NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS versions (
          id TEXT PRIMARY KEY, document_id TEXT NOT NULL, label TEXT NOT NULL,
          created_at TEXT NOT NULL, sha256 TEXT NOT NULL, source TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS save_requests (
          id TEXT PRIMARY KEY, document_id TEXT NOT NULL, state TEXT NOT NULL,
          version_id TEXT, error TEXT, created_at TEXT NOT NULL, label TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS results (
          id INTEGER PRIMARY KEY AUTOINCREMENT, document_id TEXT,
          test_id TEXT NOT NULL, status TEXT NOT NULL, details TEXT NOT NULL,
          created_at TEXT NOT NULL, evidence_type TEXT NOT NULL
        );
        """)
    asset = DATA / "assets" / "certificate.png"
    if not asset.exists():
        asset.write_bytes(certificate_png())
    with db() as con:
        empty = con.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0
    if empty:
        create_document("投标验证样本.docx", sample_docx())
    from workflow import init_workflow
    init_workflow(__import__(__name__))


@asynccontextmanager
async def lifespan(app):
    init_data()
    yield


app = FastAPI(title="投标助手文档验证台", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
                   allow_methods=["GET", "POST"], allow_headers=["Content-Type"])


def get_document(doc_id):
    with db() as con:
        row = con.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
    if not row:
        raise HTTPException(404, "文档不存在")
    return dict(row)


def create_document(title: str, content: bytes):
    doc_id = uuid.uuid4().hex
    digest = hashlib.sha256(content).hexdigest()
    (DATA / "documents" / f"{doc_id}.docx").write_bytes(content)
    with db() as con:
        con.execute("INSERT INTO documents(id,title,doc_key,created_at,last_hash,saved_at) VALUES(?,?,?,?,?,?)",
                    (doc_id, title, uuid.uuid4().hex, now(), digest, now()))
    add_version(doc_id, content, "导入原件", "initial")
    return get_document(doc_id)


def add_version(doc_id, content, label, source):
    vid = uuid.uuid4().hex
    digest = hashlib.sha256(content).hexdigest()
    (DATA / "versions" / f"{vid}.docx").write_bytes(content)
    with db() as con:
        con.execute("INSERT INTO versions VALUES(?,?,?,?,?,?)", (vid, doc_id, label, now(), digest, source))
    return vid


def ticket(kind, identifier):
    return jwt.encode({"kind": kind, "id": identifier, "exp": int(time.time()) + 86400}, APP_SECRET, algorithm="HS256")


def check_ticket(token, kind, identifier):
    try:
        data = jwt.decode(token, APP_SECRET, algorithms=["HS256"])
        if data.get("kind") != kind or data.get("id") != identifier:
            raise ValueError("mismatch")
    except (jwt.PyJWTError, ValueError):
        raise HTTPException(403, "文件或回调凭证无效或已过期") from None


def signed_url(path, kind, identifier):
    return f"{CONTAINER_BASE}{path}?ticket={ticket(kind, identifier)}"


def validate_docx(content):
    if len(content) > 40 * 1024 * 1024:
        raise HTTPException(413, "Demo 单文件上限 40 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ValueError()
            if sum(info.file_size for info in archive.infolist()) > 300 * 1024 * 1024:
                raise HTTPException(413, "DOCX 解压后内容过大")
            if any("vbaproject" in name.lower() for name in archive.namelist()):
                raise HTTPException(400, "Demo 不接受带宏的文件")
    except (zipfile.BadZipFile, ValueError):
        raise HTTPException(400, "不是有效的 DOCX 文件") from None


def safe_ds_url(url: str):
    parsed, base = urlparse(url), urlparse(DS)
    aliases = {base.hostname, "localhost", "127.0.0.1", "host.docker.internal"}
    if parsed.scheme not in {"http", "https"} or parsed.username or parsed.password:
        raise HTTPException(400, "无效的文档服务下载地址")
    if parsed.hostname not in aliases or parsed.port != base.port:
        raise HTTPException(400, "回调下载地址不属于已配置的文档服务")
    return DS + parsed.path + ("?" + parsed.query if parsed.query else "")


async def ds_post(path, payload):
    if COMMAND_ENABLED:
        payload = {**payload, "token": jwt.encode(payload, COMMAND_SECRET, algorithm="HS256")}
    try:
        async with httpx.AsyncClient(timeout=90, trust_env=False) as client:
            r = await client.post(DS + path, json=payload)
            r.raise_for_status()
            return r.json()
    except (httpx.HTTPError, ValueError):
        raise HTTPException(502, "ONLYOFFICE 服务未响应，请检查容器及服务状态") from None


@app.get("/api/health")
async def health():
    ok = False
    try:
        async with httpx.AsyncClient(timeout=4, trust_env=False) as client:
            r = await client.get(DS + "/healthcheck")
            ok = r.is_success and r.text.strip().lower() == "true"
    except httpx.HTTPError:
        pass
    return {"backend": True, "documentServer": ok, "documentServerUrl": DS,
            "containerBackendUrl": CONTAINER_BASE, "jwtEnabled": JWT_ENABLED,
            "message": "文档服务可连接" if ok else "文档服务暂不可用；检查容器内 PostgreSQL、RabbitMQ 及编辑服务"}


@app.get("/api/documents")
def documents():
    with db() as con:
        return [dict(r) for r in con.execute("SELECT * FROM documents ORDER BY created_at DESC")]


@app.post("/api/documents/sample")
def new_sample():
    return create_document("投标验证样本.docx", sample_docx())


@app.post("/api/documents/upload")
async def upload(file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".docx"):
        raise HTTPException(400, "编制工作区目前接受 DOCX；PDF 原文定位单独验证")
    content = await file.read(40 * 1024 * 1024 + 1)
    validate_docx(content)
    return create_document(Path(file.filename.replace("\\", "/")).name[:160], content)


@app.get("/api/documents/{doc_id}/config")
def editor_config(doc_id: str, user: str = "a", version: str | None = None):
    doc = get_document(doc_id)
    if user not in {"a", "b"}:
        raise HTTPException(400, "未知测试用户")
    if version:
        with db() as con:
            ver = con.execute("SELECT * FROM versions WHERE id=? AND document_id=?", (version, doc_id)).fetchone()
        if not ver:
            raise HTTPException(404, "版本不存在")
        key = "v-" + version
        url = signed_url(f"/files/versions/{version}", "version", version)
    else:
        key = doc["doc_key"]
        url = signed_url(f"/files/documents/{doc_id}", "document", doc_id)
    config = {
        "documentType": "word", "type": "desktop", "width": "100%", "height": "100%",
        "document": {"fileType": "docx", "key": key, "title": doc["title"], "url": url,
                     "permissions": {"edit": not bool(version), "download": True, "print": True,
                                     "comment": not bool(version), "review": not bool(version),
                                     "modifyContentControl": not bool(version)}},
        "editorConfig": {"mode": "view" if version else "edit", "lang": "zh-CN",
                         "user": {"id": f"demo-{user}", "name": "测试用户甲" if user == "a" else "测试用户乙"},
                         "coEditing": {"mode": "fast", "change": True},
                         "customization": {"autosave": True, "forcesave": True, "compactHeader": True,
                                           "integrationMode": "embed",
                                           "compactToolbar": True, "hideRightMenu": True,
                                           "help": False, "feedback": False}}
    }
    if not version:
        config["editorConfig"]["callbackUrl"] = signed_url(f"/api/onlyoffice/callback/{doc_id}/{key}", "callback", f"{doc_id}:{key}")
    if JWT_ENABLED:
        config["token"] = jwt.encode(config, JWT_SECRET, algorithm="HS256")
    return {"documentServerUrl": DS, "config": config}


@app.get("/files/documents/{doc_id}")
def source_file(doc_id: str, ticket: str):
    check_ticket(ticket, "document", doc_id)
    doc = get_document(doc_id)
    return FileResponse(DATA / "documents" / f"{doc_id}.docx", filename=doc["title"])


@app.get("/files/versions/{vid}")
def version_source(vid: str, ticket: str):
    check_ticket(ticket, "version", vid)
    with db() as con:
        row = con.execute("SELECT * FROM versions WHERE id=?", (vid,)).fetchone()
    if not row:
        raise HTTPException(404, "版本不存在")
    return FileResponse(DATA / "versions" / f"{vid}.docx", filename="version.docx")


@app.get("/api/assets")
def assets():
    return [{"id": "certificate", "name": "虚构企业资质证书", "width": 1000, "height": 670,
             "preview": "/api/assets/certificate.png",
             "url": signed_url("/files/certificate.png", "asset", "certificate")}]


@app.get("/api/assets/certificate.png")
def asset_preview():
    return FileResponse(DATA / "assets" / "certificate.png")


@app.get("/files/certificate.png")
def asset_source(ticket: str):
    check_ticket(ticket, "asset", "certificate")
    return FileResponse(DATA / "assets" / "certificate.png")


@app.post("/api/onlyoffice/callback/{doc_id}/{key}")
async def callback(doc_id: str, key: str, request: Request, ticket: str):
    check_ticket(ticket, "callback", f"{doc_id}:{key}")
    body = await request.json()
    if OUTBOX_ENABLED:
        raw = request.headers.get(JWT_HEADER, "").removeprefix("Bearer ") or body.get("token", "")
        try:
            decoded = jwt.decode(raw, OUTBOX_SECRET, algorithms=["HS256"])
            payload = decoded.get("payload", decoded)
            for field in ["key", "status", "url"]:
                if field in body and payload.get(field) != body[field]:
                    raise ValueError("payload mismatch")
        except (jwt.PyJWTError, ValueError):
            raise HTTPException(403, "回调签名验证失败") from None
    if body.get("key") != key:
        raise HTTPException(409, "回调文档 key 不匹配")
    async with locks.setdefault(doc_id, asyncio.Lock()):
        doc = get_document(doc_id)
        if doc["doc_key"] != key:
            return {"error": 0}  # Old session must never overwrite restored/new content.
        status = body.get("status")
        if status == 1:
            with db() as con:
                con.execute("UPDATE documents SET active_users=? WHERE id=?", (len(body.get("users", [])), doc_id))
        if status in {3, 7}:
            with db() as con:
                con.execute("UPDATE documents SET save_error=? WHERE id=?", (f"ONLYOFFICE 保存错误 {status}", doc_id))
                con.execute("UPDATE save_requests SET state='failed',error=? WHERE id=? AND document_id=?",
                            (f"ONLYOFFICE 保存错误 {status}", body.get("userdata", ""), doc_id))
            return {"error": 0}
        if status == 4:
            with db() as con:
                con.execute("UPDATE documents SET active_users=0,doc_key=? WHERE id=?", (uuid.uuid4().hex, doc_id))
        if status not in {2, 6}:
            return {"error": 0}
        try:
            url = safe_ds_url(body.get("url", ""))
            async with httpx.AsyncClient(timeout=90, trust_env=False) as client:
                r = await client.get(url)
                r.raise_for_status()
                content = r.content
            validate_docx(content)
            request_id = body.get("userdata", "")
            with db() as con:
                pending = con.execute("SELECT * FROM save_requests WHERE id=? AND document_id=?", (request_id, doc_id)).fetchone()
            label = pending["label"] if pending else ("结束编辑自动保存" if status == 2 else "编辑器保存")
            digest = hashlib.sha256(content).hexdigest()
            vid = add_version(doc_id, content, label, "callback")
            target = DATA / "documents" / f"{doc_id}.docx"
            tmp = target.with_suffix(".tmp")
            tmp.write_bytes(content)
            tmp.replace(target)
            with db() as con:
                con.execute("UPDATE documents SET saved_at=?,last_hash=?,save_error=NULL WHERE id=?", (now(), digest, doc_id))
                if status == 2:
                    con.execute("UPDATE documents SET active_users=0,doc_key=? WHERE id=?", (uuid.uuid4().hex, doc_id))
                if pending:
                    con.execute("UPDATE save_requests SET state='saved',version_id=? WHERE id=?", (vid, request_id))
            return {"error": 0}
        except Exception as exc:
            message = exc.detail if isinstance(exc, HTTPException) else "无法下载或存储回调文件"
            with db() as con:
                con.execute("UPDATE documents SET save_error=? WHERE id=?", (message, doc_id))
                con.execute("UPDATE save_requests SET state='failed',error=? WHERE id=? AND document_id=?",
                            (message, body.get("userdata", ""), doc_id))
            return {"error": 1}


class SaveBody(BaseModel):
    label: str = Field(default="手动保存版本", max_length=80)


@app.post("/api/documents/{doc_id}/save")
async def save(doc_id: str, body: SaveBody):
    doc = get_document(doc_id)
    rid = uuid.uuid4().hex
    with db() as con:
        con.execute("INSERT INTO save_requests VALUES(?,?,?,NULL,NULL,?,?)", (rid, doc_id, "pending", now(), body.label))
    try:
        result = await ds_post("/coauthoring/CommandService.ashx", {"c": "forcesave", "key": doc["doc_key"], "userdata": rid})
        code = result.get("error")
        if code == 4:
            # No new changes at the document server. Record a snapshot of last persisted file,
            # and tell the UI explicitly: this is not proof that browser edits were synchronized.
            content = (DATA / "documents" / f"{doc_id}.docx").read_bytes()
            vid = add_version(doc_id, content, body.label + "（已落盘快照）", "snapshot")
            with db() as con:
                con.execute("UPDATE save_requests SET state='unchanged',version_id=? WHERE id=?", (vid, rid))
        elif code != 0:
            raise HTTPException(502, f"文档服务拒绝保存，错误码 {code}；请先打开文档")
    except HTTPException as exc:
        with db() as con:
            con.execute("UPDATE save_requests SET state='failed',error=? WHERE id=?", (exc.detail, rid))
        raise
    return {"id": rid, "commandError": code}


@app.get("/api/save-requests/{rid}")
def save_status(rid: str):
    with db() as con:
        row = con.execute("SELECT * FROM save_requests WHERE id=?", (rid,)).fetchone()
    if not row:
        raise HTTPException(404, "保存请求不存在")
    return dict(row)


@app.get("/api/documents/{doc_id}/versions")
def versions(doc_id: str):
    get_document(doc_id)
    with db() as con:
        return [dict(r) for r in con.execute("SELECT * FROM versions WHERE document_id=? ORDER BY created_at DESC", (doc_id,))]


@app.post("/api/documents/{doc_id}/versions/{vid}/restore-copy")
def restore_copy(doc_id: str, vid: str):
    doc = get_document(doc_id)
    with db() as con:
        ver = con.execute("SELECT * FROM versions WHERE id=? AND document_id=?", (vid, doc_id)).fetchone()
    if not ver:
        raise HTTPException(404, "版本不存在")
    return create_document("恢复副本-" + doc["title"], (DATA / "versions" / f"{vid}.docx").read_bytes())


@app.get("/api/documents/{doc_id}/download")
def download(doc_id: str):
    doc = get_document(doc_id)
    return FileResponse(DATA / "documents" / f"{doc_id}.docx", filename=doc["title"])


@app.get("/api/versions/{vid}/download")
def download_version(vid: str):
    with db() as con:
        ver = con.execute("SELECT * FROM versions WHERE id=?", (vid,)).fetchone()
    if not ver:
        raise HTTPException(404, "版本不存在")
    return FileResponse(DATA / "versions" / f"{vid}.docx", filename="历史版本.docx")


@app.post("/api/documents/{doc_id}/export/pdf")
async def export_pdf(doc_id: str, version: str):
    get_document(doc_id)
    with db() as con:
        ver = con.execute("SELECT * FROM versions WHERE id=? AND document_id=?", (version, doc_id)).fetchone()
    if not ver:
        raise HTTPException(404, "请先保存并指定导出版本")
    result = await ds_post("/converter", {"async": False, "filetype": "docx", "outputtype": "pdf",
        "key": "pdf-" + version, "title": "投标验证文档", "url": signed_url(f"/files/versions/{version}", "version", version)})
    if result.get("error") or not result.get("endConvert"):
        raise HTTPException(502, f"PDF 转换未完成，错误码 {result.get('error', 'unknown')}")
    url = safe_ds_url(result.get("fileUrl", ""))
    async with httpx.AsyncClient(timeout=90, trust_env=False) as client:
        r = await client.get(url)
        r.raise_for_status()
    if not r.content.startswith(b"%PDF"):
        raise HTTPException(502, "转换结果不是 PDF")
    fid = uuid.uuid4().hex
    (DATA / "exports" / f"{fid}.pdf").write_bytes(r.content)
    return {"url": f"/api/exports/{fid}.pdf", "version": version, "size": len(r.content)}


@app.get("/api/exports/{filename}")
def export_file(filename: str):
    if not filename.endswith(".pdf") or len(filename) != 36:
        raise HTTPException(404)
    try:
        uuid.UUID(hex=filename[:-4])
    except ValueError:
        raise HTTPException(404) from None
    path = DATA / "exports" / filename
    if not path.exists():
        raise HTTPException(404)
    return FileResponse(path, filename="投标验证文档.pdf")


class ResultBody(BaseModel):
    document_id: str | None = None
    test_id: str = Field(max_length=80)
    status: str
    details: str = Field(max_length=8000)
    evidence_type: str = "manual"


@app.post("/api/results")
def record_result(body: ResultBody):
    if body.status not in {"passed", "partial", "failed", "untested", "api_ok"}:
        raise HTTPException(400, "未知结果状态")
    if body.evidence_type not in {"manual", "api", "system"}:
        raise HTTPException(400, "未知证据类型")
    if body.document_id:
        get_document(body.document_id)
    with db() as con:
        cur = con.execute("INSERT INTO results(document_id,test_id,status,details,created_at,evidence_type) VALUES(?,?,?,?,?,?)",
                          (body.document_id, body.test_id, body.status, body.details, now(), body.evidence_type))
    return {"id": cur.lastrowid}


@app.get("/api/results")
def results(document_id: str | None = None):
    with db() as con:
        if document_id:
            rows = con.execute("SELECT * FROM results WHERE document_id=? ORDER BY id DESC LIMIT 150", (document_id,))
        else:
            rows = con.execute("SELECT * FROM results ORDER BY id DESC LIMIT 150")
        return [dict(r) for r in rows]


@app.get("/api/results/report.docx")
def report():
    doc = Document()
    doc.add_heading("ONLYOFFICE 功能验证记录", 0)
    doc.add_paragraph("环境目标：9.2.1。接口返回成功不代表版式、持久化和协同行为已通过人工验证。")
    for row in results()[:100]:
        doc.add_heading(f"{row['test_id']} · {row['status']}", 2)
        doc.add_paragraph(f"{row['created_at']} / {row['evidence_type']}")
        doc.add_paragraph(row["details"])
    out = io.BytesIO()
    doc.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": 'attachment; filename="verification-report.docx"'})


@app.get("/api/documents/{doc_id}/attachments.zip")
def attachments(doc_id: str):
    doc = get_document(doc_id)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.write(DATA / "documents" / f"{doc_id}.docx", "document.docx")
        archive.write(DATA / "assets" / "certificate.png", "demo-certificate.png")
        archive.writestr("README.txt", "Demo bundle: last persisted DOCX and fictional certificate. "
                         "This is not automatic extraction of attachments from arbitrary documents.")
    return Response(out.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": 'attachment; filename="demo-bundle.zip"'})


from workflow import install_workflow
install_workflow(app, __import__(__name__))
