"""Business workflow demo. Uses the existing ONLYOFFICE storage/callback service."""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
import re
import uuid
import zipfile

import httpx
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel

from workflow_documents import (CATALOG, CATALOG_BY_KEY, fill_template, generate_template,
                                guess_fields, inspect_controls, replacement_span, repair_template_placeholders, parse_source, extract_outline, rule_analysis, sample_tender)


def init_workflow(host):
    (host.DATA / "workflow-sources").mkdir(exist_ok=True)
    with host.db() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS workflow_projects (
          id TEXT PRIMARY KEY, title TEXT NOT NULL, created_at TEXT NOT NULL, state TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS workflow_profiles (
          id TEXT PRIMARY KEY, title TEXT NOT NULL, data TEXT NOT NULL
        );
        """)
        data = {"bidder.name": "江西示范工程技术有限公司（虚构）", "bidder.legal_representative": "张示例",
                "bidder.address": "江西省南昌市示范路 100 号（虚构）", "bidder.phone": "0791-00000000",
                "project.name": "示范风电工程设计服务项目", "project.number": "DEMO-2026-001",
                "agent.name": "李示例", "agent.id_number": "DEMO-ID-0001（非真实证件）", "sign.date": "2026年09月24日",
                "finance.2023.revenue": "12,580.00", "finance.2023.profit": "1,260.00", "finance.2023.assets": "18,600.00",
                "finance.2024.revenue": "14,320.00", "finance.2024.profit": "1,480.00", "finance.2024.assets": "20,900.00"}
        con.execute("INSERT OR IGNORE INTO workflow_profiles VALUES(?,?,?)", ("demo-company", "示范企业 · 虚构数据", json.dumps(data, ensure_ascii=False)))


class AnalyzeBody(BaseModel):
    mode: str = "rules"


class PlanBody(BaseModel):
    templates: list[dict]


class ReviewBody(BaseModel):
    version_id: str
    mappings: dict[str, str]


class FillBody(BaseModel):
    profile_id: str = "demo-company"


class CheckBody(BaseModel):
    template_id: str
    version_id: str


class ProfileBody(BaseModel):
    values: dict[str, str]


def install_workflow(app, host):
    router = APIRouter(prefix="/api/workflow")
    project_locks = {}

    def get(pid):
        with host.db() as con:
            row = con.execute("SELECT state FROM workflow_projects WHERE id=?", (pid,)).fetchone()
        if not row:
            raise HTTPException(404, "项目不存在")
        state = json.loads(row["state"])
        if state.get("parser_version") != 2:
            state["blocks"] = parse_source((host.DATA / "workflow-sources" / f"{pid}.docx").read_bytes())
            state["outline"] = extract_outline(state["blocks"])
            state["parser_version"] = 2
            put(state)
        return state

    def put(state):
        with host.db() as con:
            con.execute("UPDATE workflow_projects SET state=? WHERE id=?", (json.dumps(state, ensure_ascii=False), state["id"]))
        return state

    def new_project(title, data):
        try:
            blocks = parse_source(data)
        except Exception:
            raise HTTPException(400, "无法解析 DOCX 内容，请确认文件未加密且可以正常打开") from None
        if not blocks or len(blocks) > 5000:
            raise HTTPException(400, "仅接受有正文且不超过 5000 个段落/表格块的 DOCX")
        if sum(len(b["text"]) for b in blocks) > 500000:
            raise HTTPException(400, "演示版支持最多 50 万字的 DOCX")
        pid = uuid.uuid4().hex
        source_doc = host.create_document(title, data)
        (host.DATA / "workflow-sources" / f"{pid}.docx").write_bytes(data)
        state = {"id": pid, "title": title, "created_at": host.now(), "source_doc_id": source_doc["id"],
                 "blocks": blocks, "outline": extract_outline(blocks), "parser_version": 2, "analysis": None, "templates": [], "fill_run": None, "report": {}}
        with host.db() as con:
            con.execute("INSERT INTO workflow_projects VALUES(?,?,?,?)", (pid, title, state["created_at"], json.dumps(state, ensure_ascii=False)))
        return state

    def source(state):
        return (host.DATA / "workflow-sources" / f"{state['id']}.docx").read_bytes()

    def version_data(doc_id, vid):
        with host.db() as con:
            row = con.execute("SELECT * FROM versions WHERE document_id=? AND id=?", (doc_id, vid)).fetchone()
        if not row:
            raise HTTPException(400, "所选保存版本不属于该文档")
        return (host.DATA / "versions" / f"{vid}.docx").read_bytes()

    def profile(profile_id):
        with host.db() as con:
            row = con.execute("SELECT * FROM workflow_profiles WHERE id=?", (profile_id,)).fetchone()
        if not row:
            raise HTTPException(404, "企业数据档案不存在")
        return {"id": row["id"], "title": row["title"], "values": json.loads(row["data"])}

    def validate_plan(state, candidates):
        if not candidates:
            raise HTTPException(400, "请至少添加一份模板，并确认原文范围")
        blocks, seen_blocks = state["blocks"], set()
        result = []
        for candidate in candidates:
            try:
                start, end = int(candidate["start"]), int(candidate["end"])
                title = str(candidate["title"]).strip()[:100]
                if not title or not (0 <= start <= end < len(blocks)):
                    raise ValueError()
                if seen_blocks.intersection(range(start, end + 1)):
                    raise ValueError("模板原文范围重叠，请调整起止位置")
                seen_blocks.update(range(start, end + 1))
                fields, occupied = [], {}
                for raw in candidate.get("fields", []):
                    bi, pi, a, b = [int(raw[n]) for n in ("block", "paragraph", "start", "end")]
                    if not start <= bi <= end or pi < 0:
                        raise ValueError("字段不在所选模板范围内")
                    text = blocks[bi]["paragraphs"][pi]["text"]
                    kind = raw.get("kind", "text")
                    if kind == "empty_cell":
                        if (a, b, raw["anchor"]) != (0, 0, "") or not blocks[bi]["paragraphs"][pi].get("empty_cell"):
                            raise ValueError("字段不是可填写的空白单元格")
                    elif kind != "text" or not (0 <= a < b <= len(text)) or text[a:b] != raw["anchor"]:
                        raise ValueError("字段锚点与原文不一致，请重新识别")
                    if kind == "text":
                        a, b = replacement_span(text, a, b, str(raw.get("label", "")), str(raw.get("fieldKey", "")))
                    spans = occupied.setdefault((bi, pi), [])
                    if any((a < y and b > x) or (a == b == x == y) for x, y in spans):
                        raise ValueError("同一个填写位置被重复标记")
                    spans.append((a, b))
                    key = str(raw.get("fieldKey", ""))
                    if key and key not in CATALOG_BY_KEY:
                        raise ValueError("数据库字段不在字段字典内")
                    fid = uuid.uuid4().hex[:12]
                    fields.append({"id": fid, "tag": "field:" + fid, "label": str(raw.get("label", "待填写"))[:60],
                                   "block": bi, "paragraph": pi, "start": a, "end": b, "anchor": text[a:b],
                                   "fieldKey": key, "kind": kind, "required": bool(raw.get("required", True))})
                if len(fields) > 500:
                    raise ValueError("单份模板字段超过 500 个")
                result.append({"id": uuid.uuid4().hex[:12], "title": title, "start": start, "end": end,
                               "reason": str(candidate.get("reason", "人工确认的模板范围"))[:500], "fields": fields})
            except (KeyError, TypeError, ValueError, IndexError) as exc:
                raise HTTPException(400, str(exc) or "模板范围或字段配置无效") from None
        return result

    @router.get("/settings")
    def settings():
        configured = bool(os.getenv("BID_AI_BASE_URL") and os.getenv("BID_AI_MODEL") and os.getenv("BID_AI_API_KEY"))
        return {"ai_configured": configured, "ai_model": os.getenv("BID_AI_MODEL", ""), "catalog": CATALOG,
                "profiles": [profile("demo-company")], "supported_uploads": ["docx"]}

    @router.get("/projects")
    def projects():
        with host.db() as con:
            rows = con.execute("SELECT id,title,created_at FROM workflow_projects ORDER BY created_at DESC LIMIT 50").fetchall()
        return [dict(r) for r in rows]

    @router.get("/projects/{pid}")
    def project(pid: str):
        return get(pid)

    @router.post("/sample")
    def sample():
        return new_project("示范工程招标文件.docx", sample_tender())

    @router.get("/sample.docx")
    def sample_download():
        return Response(sample_tender(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        headers={"Content-Disposition": 'attachment; filename="sample-tender.docx"'})

    @router.post("/upload")
    async def upload(file: UploadFile = File(...)):
        if not (file.filename or "").lower().endswith(".docx"):
            raise HTTPException(400, "本版仅支持 DOCX，PDF / 扫描件需接入 OCR 与版面重建后再使用")
        data = await file.read(40 * 1024 * 1024 + 1)
        host.validate_docx(data)
        title = file.filename.replace("\\", "/").split("/")[-1][:160]
        return new_project(title, data)

    @router.post("/profiles/{profile_id}")
    def update_profile(profile_id: str, body: ProfileBody):
        p = profile(profile_id)
        if any(k not in CATALOG_BY_KEY or len(v) > 2000 for k, v in body.values.items()):
            raise HTTPException(400, "未知字段或字段内容超过 2000 字")
        p["values"].update(body.values)
        with host.db() as con:
            con.execute("UPDATE workflow_profiles SET data=? WHERE id=?", (json.dumps(p["values"], ensure_ascii=False), profile_id))
        return p

    async def ai_analysis(state):
        if not settings()["ai_configured"]:
            raise HTTPException(400, "请先在 backend/.env 配置 BID_AI_BASE_URL、BID_AI_MODEL、BID_AI_API_KEY，并重启后端")
        from workflow_ai import analyze_chapters
        def record_response(result):
            folder = host.DATA / "workflow-ai-diagnostics"
            folder.mkdir(exist_ok=True)
            (folder / f"{state['id']}.json").write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
        return await analyze_chapters(state, validate_plan, record_response)

    @router.post("/projects/{pid}/analyze")
    async def analyze(pid: str, body: AnalyzeBody):
        async with project_locks.setdefault(pid, asyncio.Lock()):
            state = get(pid)
            if state["templates"]:
                raise HTTPException(409, "模板已生成。如需重新拆分，请重新上传创建项目，已有审核和成果会保留")
            if body.mode not in {"rules", "ai"}:
                raise HTTPException(400, "未知分析模式")
            analysis = await ai_analysis(state) if body.mode == "ai" else rule_analysis(state["blocks"])
            analysis.update(mode=body.mode, analyzed_at=host.now())
            state["analysis"] = analysis
            return put(state)

    @router.post("/projects/{pid}/detect-fields")
    def detect(pid: str, body: PlanBody):
        state = get(pid)
        templates = []
        for spec in body.templates:
            try:
                a, b = int(spec["start"]), int(spec["end"])
                if not 0 <= a <= b < len(state["blocks"]):
                    raise ValueError()
            except (KeyError, ValueError, TypeError):
                raise HTTPException(400, "无效的原文范围") from None
            templates.append({**spec, "fields": guess_fields(state["blocks"], a, b)})
        return {"templates": templates}

    @router.post("/projects/{pid}/generate")
    async def generate(pid: str, body: PlanBody):
        async with project_locks.setdefault(pid, asyncio.Lock()):
            state = get(pid)
            if state["templates"]:
                raise HTTPException(409, "该项目已生成模板，请使用已生成结果")
            if not state["analysis"]:
                raise HTTPException(409, "请先进行原文分析")
            specs = validate_plan(state, body.templates)
            try:
                generated = [(s, generate_template(source(state), s)) for s in specs]
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from None
            for spec, data in generated:
                d = host.create_document(spec["title"] + ".docx", data)
                spec.update(document_id=d["id"], reviewed=False, review_version_id=None, output_id=None)
            state["templates"] = specs
            return put(state)

    @router.post("/projects/{pid}/templates/{tid}/review")
    async def review(pid: str, tid: str, body: ReviewBody):
        async with project_locks.setdefault(pid, asyncio.Lock()):
            state = get(pid)
            t = next((t for t in state["templates"] if t["id"] == tid), None)
            if not t:
                raise HTTPException(404, "模板不存在")
            if state["fill_run"]:
                raise HTTPException(409, "已生成填写成果，当前模板审核版本已冻结")
            data = version_data(t["document_id"], body.version_id)
            found = inspect_controls(data)
            expected = {f["tag"] for f in t["fields"]}
            if set(body.mappings) != expected:
                raise HTTPException(400, "字段映射不完整")
            if any(len(found.get(tag, [])) != 1 for tag in expected):
                raise HTTPException(409, "存在已删除或重复的字段控件，请恢复后重新保存审核")
            for f in t["fields"]:
                key = body.mappings[f["tag"]]
                if key and key not in CATALOG_BY_KEY:
                    raise HTTPException(400, "未知数据库字段")
                f["fieldKey"] = key
            t.update(reviewed=True, review_version_id=body.version_id, reviewed_at=host.now())
            return put(state)

    @router.post("/projects/{pid}/templates/{tid}/repair-placeholders")
    async def repair_placeholders(pid: str, tid: str, body: ReviewBody):
        async with project_locks.setdefault(pid, asyncio.Lock()):
            state = get(pid)
            t = next((t for t in state["templates"] if t["id"] == tid), None)
            if not t or state["fill_run"]:
                raise HTTPException(409, "仅可修复尚未回填的模板")
            data = version_data(t["document_id"], body.version_id)
            try:
                repaired = repair_template_placeholders(data, t["fields"])
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from None
            found = inspect_controls(repaired)
            if any(len(found.get(f["tag"], [])) != 1 for f in t["fields"]):
                raise HTTPException(409, "字段控件有缺失或重复，请恢复后重试")
            if set(body.mappings) != {f["tag"] for f in t["fields"]} or any(k and k not in CATALOG_BY_KEY for k in body.mappings.values()):
                raise HTTPException(400, "字段映射无效")
            for f in t["fields"]:
                f["fieldKey"] = body.mappings[f["tag"]]
            document = host.create_document(t["title"] + "-占位修复.docx", repaired)
            t.update(document_id=document["id"], reviewed=False, review_version_id=None)
            return put(state)

    @router.get("/projects/{pid}/matches")
    def matches(pid: str, profile_id: str = "demo-company"):
        state, p = get(pid), profile(profile_id)
        rows = []
        for t in state["templates"]:
            for f in t["fields"]:
                definition = CATALOG_BY_KEY.get(f["fieldKey"], {})
                value = p["values"].get(f["fieldKey"], "")
                rows.append({"template_id": t["id"], "template_title": t["title"], **f, "value": value,
                             "source": definition.get("source", "未映射"),
                             "status": "matched" if value else "missing"})
        return {"profile": p, "rows": rows}

    @router.post("/projects/{pid}/templates/{tid}/reopen")
    async def reopen(pid: str, tid: str):
        async with project_locks.setdefault(pid, asyncio.Lock()):
            state = get(pid)
            t = next((t for t in state["templates"] if t["id"] == tid), None)
            if not t or not t["reviewed"] or state["fill_run"]:
                raise HTTPException(409, "仅可在首次回填之前重新编辑已审核模板")
            data = version_data(t["document_id"], t["review_version_id"])
            document = host.create_document(t["title"] + "-重新审核.docx", data)
            t.update(document_id=document["id"], reviewed=False, review_version_id=None)
            return put(state)

    @router.post("/projects/{pid}/fill")
    async def fill(pid: str, body: FillBody):
        async with project_locks.setdefault(pid, asyncio.Lock()):
            state = get(pid)
            if not state["templates"] or not all(t["reviewed"] for t in state["templates"]):
                raise HTTPException(409, "请先逐份审核模板")
            rows = matches(pid, body.profile_id)["rows"]
            values = {r["tag"]: r["value"] for r in rows}
            created = []
            try:
                for t in state["templates"]:
                    data = version_data(t["document_id"], t["review_version_id"])
                    generated = fill_template(data, values)
                    created.append((t, generated))
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from None
            for t, data in created:
                d = host.create_document(t["title"] + "-已回填.docx", data)
                t["output_id"] = d["id"]
            state["fill_run"] = {"id": uuid.uuid4().hex, "created_at": host.now(), "profile_id": body.profile_id,
                                  "rows": rows, "matched": sum(r["status"] == "matched" for r in rows),
                                  "missing": sum(r["status"] == "missing" for r in rows)}
            state["report"] = {}
            return put(state)

    @router.post("/projects/{pid}/check")
    async def check(pid: str, body: CheckBody):
        async with project_locks.setdefault(pid, asyncio.Lock()):
            state = get(pid)
            t = next((t for t in state["templates"] if t["id"] == body.template_id), None)
            if not t or not t["output_id"]:
                raise HTTPException(400, "请先生成填写成果")
            data = version_data(t["output_id"], body.version_id)
            found, issues = inspect_controls(data), []
            expectations = {r["tag"]: r for r in state["fill_run"]["rows"]}
            for f in t["fields"]:
                texts = found.get(f["tag"], [])
                if len(texts) != 1:
                    issues.append({"tag": f["tag"], "label": f["label"], "level": "error", "message": "字段控件缺失或重复"})
                    continue
                value = texts[0].strip()
                if f["required"] and (not value or value == "【" + f["label"] + "】"):
                    issues.append({"tag": f["tag"], "label": f["label"], "level": "error", "message": "必填内容尚未补充"})
                elif expectations[f["tag"]]["value"] and value != str(expectations[f["tag"]]["value"]):
                    issues.append({"tag": f["tag"], "label": f["label"], "level": "warning", "message": "当前文字与回填时的数据快照不同，请核实人工修改"})
            from docx import Document
            from docx.oxml.ns import qn
            original = Document(io.BytesIO(version_data(t["document_id"], t["review_version_id"])))
            actual = Document(io.BytesIO(data))
            # Compare fixed text independently of generated field values.
            def fixed_text(doc):
                root = doc.element.body
                for sdt in list(root.iter(qn("w:sdt"))):
                    sdt.getparent().remove(sdt)
                return "".join(n.text or "" for n in root.iter(qn("w:t")))
            if fixed_text(original) != fixed_text(actual):
                issues.append({"tag": "", "label": "固定文字", "level": "warning", "message": "模板固定文字发生变化，请对照审核版本确认"})
            report = {"version_id": body.version_id, "document_id": t["output_id"], "checked_at": host.now(),
                      "sha256": hashlib.sha256(data).hexdigest(), "field_count": len(t["fields"]), "issues": issues,
                      "passed": not any(i["level"] == "error" for i in issues),
                      "visual_review": "待人工核对字体、分页、合并单元格及招标格式要求；自动检查仅覆盖字段和固定文字。"}
            state["report"][t["id"]] = report
            put(state)
            return report

    @router.get("/projects/{pid}/bundle.zip")
    def bundle(pid: str):
        state = get(pid)
        if not state["templates"] or any(t["id"] not in state["report"] for t in state["templates"]):
            raise HTTPException(409, "请逐份保存并检查成果后再打包")
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for i, t in enumerate(state["templates"], 1):
                report = state["report"][t["id"]]
                if not report["passed"]:
                    raise HTTPException(409, "存在未解决的必填字段或控件错误，请修正并重新检查")
                safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", t["title"])
                z.writestr(f"{i:02d}-{safe_name}.docx", version_data(t["output_id"], report["version_id"]))
            z.writestr("字段与检查清单.json", json.dumps({"project": state["title"], "fill_run": state["fill_run"],
                "templates": state["templates"], "report": state["report"]}, ensure_ascii=False, indent=2))
            z.writestr("说明.txt", "本包使用逐份检查时指定的已保存版本，不包含检查之后的编辑。演示数据均为虚构。自动检查不代表格式合规，须人工核对版式。")
        return Response(out.getvalue(), media_type="application/zip", headers={"Content-Disposition": 'attachment; filename="bid-templates.zip"'})

    app.include_router(router)
